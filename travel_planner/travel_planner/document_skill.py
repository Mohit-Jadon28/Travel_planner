
from __future__ import annotations

import os
from typing import Callable

from .config import COMPRESSION, GUARDRAILS
from .guardrails import sanitize_free_text


def _read_txt_or_md(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


def _read_pdf(path: str) -> str:
    from pypdf import PdfReader
    reader = PdfReader(path)
    return "\n".join((p.extract_text() or "") for p in reader.pages[:15])


def _read_docx(path: str) -> str:
    import docx
    d = docx.Document(path)
    return "\n".join(p.text for p in d.paragraphs)


def _read_xlsx(path: str) -> str:
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    lines = []
    for ws in wb.worksheets[:3]:
        lines.append(f"[Sheet: {ws.title}]")
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= 40:
                lines.append("...")
                break
            lines.append(", ".join("" if c is None else str(c) for c in row))
    return "\n".join(lines)


_READERS: dict[str, Callable[[str], str]] = {
    ".md": _read_txt_or_md, ".txt": _read_txt_or_md,
    ".pdf": _read_pdf, ".docx": _read_docx,
    ".xlsx": _read_xlsx, ".xlsm": _read_xlsx,
    ".csv": _read_txt_or_md,
}


class DocumentSkill:
    def __init__(self, max_chars_per_file: int | None = None):
        self.max_chars_per_file = max_chars_per_file or COMPRESSION["skill_excerpt_chars"]

    def read_one(self, path: str) -> str:
        ext = os.path.splitext(path)[1].lower()
        print(f"   \U0001F4C2 Skill: Document Ingestion accessing '{path}' (.{ext.lstrip('.')})...")
        reader = _READERS.get(ext)
        if reader is None:
            return f"[skipped {path}: unsupported extension '{ext}']"
        try:
            size = os.path.getsize(path)
            if size > GUARDRAILS["max_reference_file_bytes"]:
                return f"[skipped {path}: file too large ({size} bytes)]"
            raw = reader(path)
            clean = sanitize_free_text(raw, self.max_chars_per_file)
            return f"[{os.path.basename(path)}]\n{clean}"
        except Exception as e:  # noqa: BLE001
            return f"[error reading {path}: {e}]"

    def read_all(self, paths: list[str]) -> str:
        if not paths:
            return ""
        excerpts = [self.read_one(p) for p in paths]
        return "\n\n".join(excerpts)


document_skill = DocumentSkill()
